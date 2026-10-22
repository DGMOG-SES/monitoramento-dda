#!/usr/bin/env python3
"""
Extração automática do painel "Monitoramento das DDA" (Tableau Public)
para o formato das abas aba12, aba15, aba16, aba17, aba23, aba24, aba25, aba26
(e municipiosAlerta) do arquivo "Consolidado - DGMOG.xlsx".

Como funciona
-------------
Para cada aba, o script abre a view correspondente do painel com a biblioteca
tableauscraper (mesma sessão VizQL que o navegador usa), aplica o filtro de UF
e pede ao servidor o mesmo "Download > Tabela de referência cruzada" (crosstab)
que hoje é feito manualmente. O crosstab chega como TSV e é lido com pandas,
o que reproduz exatamente o layout atual das abas (inclusive os cabeçalhos
"Unnamed: 4" e as células vazias gravadas como "NaN").

Se o crosstab falhar, há um plano B: baixar os "dados de resumo" (formato
longo) e pivotar localmente.

Uso
---
    pip install tableauscraper pandas openpyxl

    # 1) Diagnóstico: lista planilhas, filtros e salva os crosstabs brutos
    python extrair_dda.py --diagnostico

    # 2) Extração: gera um xlsx novo só com as abas
    python extrair_dda.py --saida Consolidado_DDA.xlsx

    # 3) Extração comparando com o arquivo atual (confere cabeçalhos/linhas)
    python extrair_dda.py --saida Consolidado_DDA.xlsx --modelo "Consolidado - DGMOG.xlsx"

    # 4) Atualizar as abas dentro de uma cópia do arquivo atual
    #    (mantém etiologia, mortalidade etc. intactas)
    python extrair_dda.py --modelo "Consolidado - DGMOG.xlsx" --atualizar-modelo

    # Só algumas abas
    python extrair_dda.py --abas aba23,aba24
"""

import argparse
import json
import logging
import re
import shutil
import sys
import time
import unicodedata
from datetime import datetime
from io import StringIO
from pathlib import Path

import pandas as pd
import requests

try:
    from tableauscraper import TableauScraper
    from tableauscraper import api as ts_api
    from tableauscraper import utils as ts_utils
except ImportError:  # pragma: no cover
    sys.exit("Instale as dependências: pip install tableauscraper pandas openpyxl")

BASE = "https://public.tableau.com/views/MonitoramentodasDDA"
UF_SIGLA = "PE"
UF_NOME = "Pernambuco"

# Mapeamento aba -> view do painel. "filtros" aceita vários nomes candidatos
# para o mesmo campo, porque o nome exibido do filtro pode variar entre views.
# Valor "*" significa "selecionar todos os valores" (útil para Ano).
# O painel abre filtrado em AC e só nos últimos anos: é preciso trocar a UF
# e marcar todos os anos.
ANO = ("Ano - Divisão (Grupo C)", "Ano")

ABAS = {
    "aba12": {
        "view": "12-CasosSEMun_1",
        "titulo": "12-Casos /SE (Mun)",
        "filtros": {("Uf Sigla", "UF_Sigla", "UF", "Uf"): [UF_SIGLA, UF_NOME], ANO: "*"},
    },
    "aba15": {
        "view": "15-CasosFxEtSEeAnoMun",
        "titulo": "15-Casos Fx Et/SE e Ano (Mun)",
        "filtros": {("Uf Sigla", "UF_Sigla", "UF", "Uf"): [UF_SIGLA, UF_NOME], ANO: "*"},
    },
    "aba16": {
        "view": "16-CasosPlanTratSEeAnoMun",
        "titulo": "16-Casos Plan Trat /SE e Ano (Mun)",
        "filtros": {("Uf Sigla", "UF_Sigla", "UF", "Uf"): [UF_SIGLA, UF_NOME], ANO: "*"},
    },
    "aba17": {
        "view": "17-US-DDASEMun",
        "titulo": "17-US-DDA/SE (Mun)",
        "filtros": {("Uf Sigla", "UF_Sigla", "UF", "Uf"): [UF_SIGLA, UF_NOME], ANO: "*"},
    },
    "aba23": {
        "view": "23-LimitesDC2026UF",
        "titulo": "23-Limites DC/2026 (UF)",
        "filtros": {("UF_Sigla", "Uf Sigla", "UF", "Uf"): [UF_NOME, UF_SIGLA]},
    },
    "aba24": {
        "view": "24-LimitesDC2025UF",
        "titulo": "24-Limites DC/2025 (UF)",
        "filtros": {("UF_Sigla", "Uf Sigla", "UF", "Uf"): [UF_NOME, UF_SIGLA]},
    },
    "aba25": {
        "view": "25-LimitesDC2026Mun",
        "titulo": "25-Limites DC/2026 (Mun)",
        "filtros": {("Uf Sigla", "UF_Sigla", "UF", "Uf"): [UF_SIGLA, UF_NOME]},
    },
    "aba26": {
        "view": "26-LimitesDC2025Mun",
        "titulo": "26-Limites DC/2025 (Mun)",
        "filtros": {("Uf Sigla", "UF_Sigla", "UF", "Uf"): [UF_SIGLA, UF_NOME]},
    },
}

# municipiosAlerta é derivada de aba25 (2026) e aba26 (2025): formato longo.
ALERTA_FONTES = {"aba25": 2026, "aba26": 2025}

log = logging.getLogger("dda")
REPETIR_ROTULOS = False  # --repetir-rotulos: mantém UF/município/ano em todas as linhas


# --------------------------------------------------------------------------
# utilidades
# --------------------------------------------------------------------------
def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


def presmodel_holder(resp):
    return resp["vqlCmdResponse"]["layoutStatus"]["applicationPresModel"][
        "presentationLayerNotification"][0]["presModelHolder"]


def para_numero(v):
    """Converte '1.126', '1,126', '-723', '12,5' em número; mantém texto."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return v
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(" ", "").replace(" ", "")
    if s == "":
        return None
    if re.fullmatch(r"-?\d{1,3}([.,]\d{3})+", s):  # separador de milhar
        return float(re.sub(r"[.,]", "", s))
    if re.fullmatch(r"-?\d+([.,]\d+)?", s):
        return float(s.replace(",", "."))
    return v


# --------------------------------------------------------------------------
# acesso ao Tableau
# --------------------------------------------------------------------------
HOST = "https://public.tableau.com"
WB = "MonitoramentodasDDA"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/129.0 Safari/537.36")


class SessaoTableau(requests.Session):
    """Repete pedidos que voltam 410 (Gone), 5xx ou vazios, com espera crescente.
    O roteamento ao servidor certo é feito pelo cabeçalho global-session-header
    (ver abrir_view); as repetições cobrem falhas pontuais e o tempo de preparo
    do arquivo do crosstab."""

    REPETIR = {410, 429, 500, 502, 503, 504}
    historico = []  # (endpoint, status, bytes) de cada tentativa, para o diagnóstico

    def request(self, method, url, *a, **kw):
        kw.setdefault("timeout", 180)
        fim = url.split("?")[0].rstrip("/").split("/")[-1] or url.split("/")[-2]
        espera = 0.5
        for i in range(10):
            r = super().request(method, url, *a, **kw)
            vazio = len(r.content) == 0 and ("/commands/" in url or "/tempfile/" in url
                                              or "bootstrapSession" in url)
            SessaoTableau.historico.append((fim, r.status_code, len(r.content)))
            if r.status_code not in self.REPETIR and not vazio:
                return r
            log.debug("%s HTTP %s (%d bytes), tentativa %d; aguardando %.1fs",
                      fim, r.status_code, len(r.content), i + 1, espera)
            time.sleep(espera)
            espera = min(espera * 2, 20)
        return r


def abrir_view(view):
    """Fluxo novo do Tableau Public: a página não traz mais o tsConfigContainer
    preenchido; a configuração da sessão vem de POST .../startSession/viewing."""
    ts = TableauScraper(logLevel=logging.WARNING, delayMs=300)
    s = SessaoTableau()
    s.headers.update({"User-Agent": UA})
    ts.session = s
    pagina = f"{BASE}/{view}"
    log.info("abrindo %s", pagina)
    s.get(pagina, params={":embed": "y", ":showVizHome": "no"})
    r = s.post(f"{HOST}/vizql/w/{WB}/v/{view}/startSession/viewing",
               params={":embed": "y", ":showVizHome": "no"},
               headers={"Accept": "application/json", "X-Requested-With": "XMLHttpRequest",
                        "Referer": pagina, "Origin": HOST})
    r.raise_for_status()
    # O cabeçalho global-session-header (base64 de "ip:porta" do servidor que criou
    # a sessão) precisa voltar em todos os pedidos seguintes; é ele que faz o
    # balanceador mandar o pedido ao servidor certo. Sem ele, quase tudo dá 410.
    gsh = r.headers.get("global-session-header")
    if gsh:
        s.headers["global-session-header"] = gsh
    else:
        log.warning("startSession sem global-session-header; pedidos podem dar 410")
    cfg = r.json()
    cfg.setdefault("vizql_root", f"/vizql/w/{WB}/v/{view}")
    cfg.setdefault("sheetId", cfg.get("sheetName") or view)
    if "sessionid" not in cfg:
        raise RuntimeError(f"startSession sem sessionid; chaves: {list(cfg)[:40]}")
    ts.tableauData = cfg
    ts.host = HOST

    form = {
        "worksheetPortSize": json.dumps({"w": 1920, "h": 1080}),
        "dashboardPortSize": json.dumps({"w": 1920, "h": 1080}),
        "clientDimension": json.dumps({"w": 1920, "h": 1080}),
        "renderMapsClientSide": "true",
        "isBrowserRendering": "true",
        "browserRenderingThreshold": "100",
        "formatDataValueLocally": "false",
        "navType": "Nav",
        "navSrc": "Top",
        "devicePixelRatio": "1",
        "clientRenderPixelLimit": "16777216",
        "allowAutogenWorksheetPhoneLayouts": "false",
        "sheet_id": cfg["sheetId"],
        "showParams": json.dumps({"checkpoint": False, "refresh": False, "refreshUnmodified": False,
                                  "unknownParams": ":embed=y&:showVizHome=no"}),
        "filterTileSize": "200",
        "locale": "pt_BR",
        "language": "pt_BR",
        "verboseMode": "false",
        ":session_feature_flags": "{}",
        "keychain_version": "1",
    }
    if cfg.get("stickySessionKey"):
        sk = cfg["stickySessionKey"]
        form["stickySessionKey"] = sk if isinstance(sk, str) else json.dumps(sk)
    url_boot = f"{HOST}{cfg['vizql_root']}/bootstrapSession/sessions/{cfg['sessionid']}"
    r = s.post(url_boot, data=form, headers={"Referer": pagina, "Origin": HOST,
                                             "X-Requested-With": "XMLHttpRequest"})
    texto = r.text
    m = re.search(r"\d+;({.*})\d+;({.*})", texto, re.S)
    if not m:
        raise RuntimeError(f"bootstrapSession HTTP {r.status_code}: {texto[:300]!r}")
    ts.lastActionTime = time.time()
    ts.info = json.loads(m.group(1))
    ts.data = json.loads(m.group(2))
    if "presModelMap" in ts.data.get("secondaryInfo", {}):
        pmm = ts.data["secondaryInfo"]["presModelMap"]
        ts.dataSegments = pmm["dataDictionary"]["presModelHolder"]["genDataDictionaryPresModel"]["dataSegments"]
        ts.parameters = ts_utils.getParameterControlInput(ts.info)
    ts.dashboard = ts.info["sheetName"]
    ts.filters = ts_utils.getFiltersForAllWorksheet(ts.logger, ts.data, ts.info, rootDashboard=ts.dashboard)
    return ts, ts.getWorkbook()


def descrever_filtros(wb):
    out = {}
    for ws in wb.worksheets:
        out[ws.name] = [
            {
                "coluna": f.get("column"),
                "globalFieldName": f.get("globalFieldName"),
                "n_valores": len(f.get("values", [])),
                "valores_exemplo": f.get("values", [])[:15],
                "selecionados": f.get("selection", [])[:15],
                "ordinal": f.get("ordinal"),
            }
            for f in ws.getFilters()
        ]
    return out


def aplicar_filtros(ts, wb, filtros):
    """Aplica cada filtro configurado no primeiro worksheet que o expõe."""
    for nomes, valores in filtros.items():
        aplicado = False
        alvos = {norm(n) for n in nomes}
        for ws in wb.worksheets:
            for f in ws.getFilters():
                if norm(f.get("column")) not in alvos:
                    continue
                dominio = f.get("values", [])
                if valores == "*":
                    escolha = list(dominio)
                else:
                    escolha = next((v for v in valores if v in dominio), None)
                    if escolha is None:
                        # tenta casar sem acento/maiúscula
                        mapa = {norm(d): d for d in dominio}
                        escolha = next((mapa[norm(v)] for v in valores if norm(v) in mapa), None)
                if escolha is None or escolha == []:
                    log.warning("filtro %s: nenhum de %s está no domínio (%s...)",
                                f.get("column"), valores, dominio[:10])
                    continue
                log.info("filtro %s = %s (worksheet %s)", f.get("column"),
                         escolha if valores != "*" else "(todos)", ws.name)
                novo = ws.setFilter(f["column"], escolha)
                if not novo.worksheets:
                    novo = ws.setFilter(f["column"], escolha, dashboardFilter=True)
                if novo.worksheets:
                    wb = novo
                aplicado = True
                break
            if aplicado:
                break
        if not aplicado:
            log.warning("filtro %s não encontrado nesta view (pode já estar fixo no painel)", nomes)
    return wb


def listar_planilhas_crosstab(ts):
    r = ts_api.exportCrosstabServerDialog(ts)
    itens = presmodel_holder(r)["genExportCrosstabOptionsDialogPresModel"]["thumbnailSheetPickerItems"]
    return [{"sheetName": i["sheetName"], "sheetdocId": i["sheetdocId"]} for i in itens]


def baixar_crosstab(ts, preferida=None):
    """Executa o mesmo 'Download > Crosstab' do navegador. Retorna (nome, DataFrame)."""
    itens = listar_planilhas_crosstab(ts)
    if not itens:
        raise RuntimeError("o painel não ofereceu nenhuma planilha para crosstab")
    escolhido = None
    if preferida:
        escolhido = next((i for i in itens if norm(i["sheetName"]) == norm(preferida)), None)
    if escolhido is None:
        if len(itens) > 1:
            log.info("planilhas disponíveis para crosstab: %s", [i["sheetName"] for i in itens])
        escolhido = itens[0]
    r = ts_api.exportCrosstabToCsvServer(ts, escolhido["sheetdocId"])
    h = presmodel_holder(r)
    if "genExportFilePresModel" in h:
        chave = h["genExportFilePresModel"]["resultKey"]
    elif "genFileDownloadPresModel" in h:
        chave = h["genFileDownloadPresModel"]["tempfileKey"]
    else:
        raise RuntimeError(f"resposta inesperada do export: {list(h)}")
    texto = ts_api.downloadCrossTabData(ts, chave)
    df = pd.read_csv(StringIO(texto), sep="\t", dtype=str, keep_default_na=False)
    return escolhido["sheetName"], df, texto


def baixar_resumo(wb, preferida=None):
    """Plano B: dados de resumo (formato longo) do worksheet principal."""
    alvos = wb.worksheets
    if preferida:
        alvos = sorted(alvos, key=lambda w: norm(w.name) != norm(preferida))
    for ws in alvos:
        try:
            d = ws.getDownloadableSummaryData(numRows=2_000_000)
            if d is not None and len(d.data):
                return ws.name, d.data
        except Exception as e:  # noqa: BLE001
            log.warning("resumo de %s falhou: %s", ws.name, e)
    raise RuntimeError("nenhum worksheet devolveu dados de resumo")


# --------------------------------------------------------------------------
# pós-processamento para o formato das abas
# --------------------------------------------------------------------------
def formatar_crosstab(df):
    """Deixa o crosstab no mesmo formato que as abas atuais."""
    df = df.copy()
    # cabeçalhos vazios viram "Unnamed: i", como no pandas.read_csv original
    df.columns = [c if str(c).strip() else f"Unnamed: {i}" for i, c in enumerate(df.columns)]
    df = df.replace({"": None})
    # colunas de semana epidemiológica: '1' -> '01'
    novas = []
    for c in df.columns:
        s = str(c).strip()
        novas.append(s.zfill(2) if re.fullmatch(r"\d{1,2}", s) else s)
    df.columns = novas
    # conversão numérica nas colunas de valores (SE e Total)
    for c in df.columns:
        if re.fullmatch(r"\d{2}", c) or c in ("Total", "NaN", "Null", "Nulo"):
            df[c] = df[c].map(para_numero)
    if "Mun IBGE" in df.columns:
        df["Mun IBGE"] = df["Mun IBGE"].map(para_numero)
    for c in df.columns:
        if "Ano" in c:
            df[c] = df[c].map(para_numero)
    return df


def ocultar_rotulos_repetidos(df):
    """O crosstab atual do Tableau repete os rótulos de linha (UF, município, ano)
    em todas as linhas; as abas antigas traziam esses rótulos só na primeira linha
    de cada grupo. Para manter o formato que o BI já lê, apaga as repetições
    (hierarquicamente: um nível só é apagado se os níveis à esquerda também se repetem)."""
    dims = []
    for c in df.columns:
        if str(c).startswith("Unnamed") or c == "Total" or re.fullmatch(r"\d{2}", str(c)):
            break
        dims.append(c)
    if not dims:
        return df
    orig = df[dims].copy()
    igual = pd.Series(True, index=df.index)
    for c in dims:
        igual &= orig[c].eq(orig[c].shift())
        df.loc[igual, c] = None
    return df


def gerar_municipios_alerta(abas):
    partes = []
    for aba, ano in ALERTA_FONTES.items():
        if aba not in abas:
            continue
        d = abas[aba].copy()
        rot = d.columns[3]
        d["Mun IBGE"] = pd.to_numeric(d["Mun IBGE"].replace("NaN", None), errors="coerce").ffill()
        d["k"] = d[rot].astype(str).str.replace(r"\s+Munic.*$|\s+Mun em rela.*$", "", regex=True)
        d = d[d["k"].isin(["Casos", "Limite Superior", "Mediana", "Limite Inferior"])]
        se = [c for c in d.columns if re.fullmatch(r"\d{2}", str(c))]
        lg = d.melt(id_vars=["Mun IBGE", "k"], value_vars=se, var_name="Semanas", value_name="v")
        lg["v"] = pd.to_numeric(lg["v"].replace("NaN", None), errors="coerce")
        p = lg.pivot_table(index=["Mun IBGE", "Semanas"], columns="k", values="v",
                           aggfunc="first", dropna=False).reset_index()
        p["Semanas"] = p["Semanas"].astype(int)
        p["Mun IBGE"] = p["Mun IBGE"].astype(int)
        p["ano"] = ano
        p = p.sort_values(["Semanas", "Mun IBGE"])
        partes.append(p[["Mun IBGE", "Semanas", "Casos", "Limite Superior",
                         "Mediana", "Limite Inferior", "ano"]])
    return pd.concat(partes, ignore_index=True) if partes else None


def comparar_com_modelo(nome, novo, modelo):
    msgs = []
    if list(map(str, novo.columns)) != list(map(str, modelo.columns)):
        so_novo = [c for c in novo.columns if c not in modelo.columns]
        so_mod = [c for c in modelo.columns if c not in novo.columns]
        msgs.append(f"colunas diferentes (só no novo: {so_novo[:10]}; só no modelo: {so_mod[:10]})")
    if abs(len(novo) - len(modelo)) > 0:
        msgs.append(f"linhas: novo={len(novo)} modelo={len(modelo)}")
    if "Total" in novo.columns and "Total" in modelo.columns:
        try:
            t1 = pd.to_numeric(novo["Total"].replace("NaN", None), errors="coerce").iloc[-1]
            t2 = pd.to_numeric(modelo["Total"].replace("NaN", None), errors="coerce").iloc[-1]
            msgs.append(f"Total geral: novo={t1:,.0f} modelo={t2:,.0f}")
        except Exception:  # noqa: BLE001
            pass
    print(f"  [{nome}] " + ("; ".join(msgs) if msgs else "estrutura idêntica ao modelo"))


# --------------------------------------------------------------------------
# orquestração
# --------------------------------------------------------------------------
def extrair_aba(nome, cfg, pasta_bruta=None, tentativas=3):
    ultimo = None
    for t in range(1, tentativas + 1):
        try:
            ts, wb = abrir_view(cfg["view"])
            wb = aplicar_filtros(ts, wb, cfg.get("filtros", {}))
            try:
                planilha, df, bruto = baixar_crosstab(ts, cfg.get("planilha"))
                if pasta_bruta:
                    (pasta_bruta / f"{nome}_crosstab.tsv").write_text(bruto, encoding="utf-8")
                log.info("%s: crosstab '%s' %s", nome, planilha, df.shape)
                df = formatar_crosstab(df)
                return df if REPETIR_ROTULOS else ocultar_rotulos_repetidos(df)
            except Exception as e:  # noqa: BLE001
                log.warning("%s: crosstab falhou (%s); tentando dados de resumo", nome, e)
                planilha, longo = baixar_resumo(wb, cfg.get("planilha"))
                if pasta_bruta:
                    longo.to_csv(pasta_bruta / f"{nome}_resumo_longo.csv", index=False)
                log.warning("%s: salvos dados em formato LONGO de '%s' %s — revise o pivot",
                            nome, planilha, longo.shape)
                return longo
        except Exception as e:  # noqa: BLE001
            ultimo = e
            log.warning("%s: tentativa %d falhou: %s", nome, t, e)
            time.sleep(3 * t)
    raise RuntimeError(f"{nome}: falhou após {tentativas} tentativas: {ultimo}")


def diagnostico(abas, pasta):
    pasta.mkdir(parents=True, exist_ok=True)
    relatorio = {"gerado_em": datetime.now().isoformat(), "views": {}}
    for nome in abas:
        cfg = ABAS[nome]
        item = {"view": cfg["view"]}
        try:
            ts, wb = abrir_view(cfg["view"])
            item["sheets"] = wb.getSheets()
            item["worksheets"] = [ws.name for ws in wb.worksheets]
            item["filtros"] = descrever_filtros(wb)
            item["parametros"] = wb.getParameters()
            try:
                item["colunas_por_worksheet"] = {ws.name: ws.getColumns() for ws in wb.worksheets}
            except Exception as e:  # noqa: BLE001  (view renderizada no servidor)
                item["colunas_por_worksheet"] = f"indisponível: {e!r}"
            wb = aplicar_filtros(ts, wb, cfg.get("filtros", {}))
            item["filtros_depois"] = descrever_filtros(wb)
            item["planilhas_crosstab"] = listar_planilhas_crosstab(ts)
            planilha, df, bruto = baixar_crosstab(ts, cfg.get("planilha"))
            (pasta / f"{nome}_crosstab.tsv").write_text(bruto, encoding="utf-8")
            item["crosstab"] = {"planilha": planilha, "shape": list(df.shape),
                                "colunas": list(df.columns)[:12],
                                "primeiras_linhas": df.head(5).values.tolist()}
        except Exception as e:  # noqa: BLE001
            import traceback
            item["erro"] = repr(e)
            item["traceback"] = traceback.format_exc()
            log.debug(item["traceback"])
        item["requisicoes"] = SessaoTableau.historico[-60:]
        SessaoTableau.historico.clear()
        relatorio["views"][nome] = item
        print(f"{nome}: {'ERRO ' + item['erro'] if 'erro' in item else 'ok'}")
    arq = pasta / "diagnostico.json"
    arq.write_text(json.dumps(relatorio, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"\nRelatório salvo em {arq}")


N_MUNICIPIOS_PE = 185


def validar(abas):
    """Confere o resultado antes de publicar, para nunca sobrescrever a planilha
    do painel com dados quebrados. Levanta erro se algo estiver fora do esperado."""
    problemas = []
    for nome in ABAS:
        if nome not in abas:
            continue
        df = abas[nome]
        se = [c for c in df.columns if re.fullmatch(r"\d{2}", str(c))]
        if len(df) == 0:
            problemas.append(f"{nome}: vazia")
        if len(se) < 52:
            problemas.append(f"{nome}: só {len(se)} colunas de semana")
        if "Mun IBGE" in df.columns:
            muns = pd.to_numeric(df["Mun IBGE"], errors="coerce").dropna()
            muns = muns[(muns >= 260000) & (muns < 270000)].nunique()
            if muns != N_MUNICIPIOS_PE:
                problemas.append(f"{nome}: {muns} municípios de PE (esperado {N_MUNICIPIOS_PE})")
        if nome in ("aba23", "aba24") and len(df) != 5:
            problemas.append(f"{nome}: {len(df)} linhas (esperado 5)")
    if "municipiosAlerta" in abas and abas["municipiosAlerta"]["Casos"].notna().sum() == 0:
        problemas.append("municipiosAlerta: sem casos")
    if problemas:
        raise RuntimeError("validação falhou:\n  " + "\n  ".join(problemas))
    log.info("validação ok: %s", ", ".join(f"{k}={len(v)}" for k, v in abas.items()))


def gravar_csv(abas, pasta, na_rep="NaN"):
    """Um CSV por aba (UTF-8), com o mesmo conteúdo que vai para a planilha."""
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    for nome, df in abas.items():
        df = df.copy()
        for c in df.columns:  # 1.0 -> 1 nas colunas de contagem
            if df[c].dtype.kind == "f" and (df[c].dropna() % 1 == 0).all():
                df[c] = df[c].astype("Int64")
        df.to_csv(pasta / f"{nome}.csv", index=False, encoding="utf-8",
                  na_rep="" if nome == "municipiosAlerta" else na_rep)
    log.info("CSVs gravados em %s", pasta)


def gravar_xlsx(abas, destino, na_rep="NaN"):
    with pd.ExcelWriter(destino, engine="openpyxl") as xw:
        for nome, df in abas.items():
            df.to_excel(xw, sheet_name=nome, index=False, na_rep=na_rep if nome != "municipiosAlerta" else "")


def atualizar_modelo(abas, modelo, na_rep="NaN"):
    """Grava uma cópia do modelo com as abas substituídas, numa única passada
    (abrir e salvar um xlsx deste tamanho leva alguns minutos)."""
    from openpyxl import load_workbook
    copia = modelo.with_name(f"{modelo.stem}_atualizado_{datetime.now():%Y%m%d_%H%M}{modelo.suffix}")
    log.info("abrindo %s (pode levar 1 a 2 minutos)...", modelo.name)
    wbx = load_workbook(modelo)
    for nome, df in abas.items():
        vazio = "" if nome == "municipiosAlerta" else na_rep
        if nome in wbx.sheetnames:
            pos = wbx.sheetnames.index(nome)
            del wbx[nome]
        else:
            pos = len(wbx.sheetnames)
        ws = wbx.create_sheet(nome, pos)
        ws.append([str(c) for c in df.columns])
        for linha in df.itertuples(index=False, name=None):
            ws.append([vazio if (v is None or (isinstance(v, float) and pd.isna(v))) else v
                       for v in linha])
        log.info("aba %s gravada (%d linhas)", nome, len(df))
    log.info("salvando %s (pode levar 1 a 2 minutos)...", copia.name)
    wbx.save(copia)
    return copia


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--abas", default=",".join(ABAS), help="lista separada por vírgula")
    ap.add_argument("--saida", default=f"Consolidado_DDA_{datetime.now():%Y%m%d}.xlsx")
    ap.add_argument("--modelo", help="arquivo Consolidado atual, para comparar/atualizar")
    ap.add_argument("--atualizar-modelo", action="store_true",
                    help="grava uma cópia do modelo com as abas substituídas")
    ap.add_argument("--diagnostico", action="store_true")
    ap.add_argument("--brutos", default="brutos_tableau", help="pasta para os arquivos brutos")
    ap.add_argument("--repetir-rotulos", action="store_true",
                    help="mantém UF/município/ano em todas as linhas (padrão: só na 1ª do grupo, como as abas atuais)")
    ap.add_argument("--csv-dir", help="grava também um CSV por aba nesta pasta")
    ap.add_argument("--sem-xlsx", action="store_true", help="não grava o .xlsx de saída")
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args()

    logging.basicConfig(level=logging.DEBUG if a.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")
    global REPETIR_ROTULOS
    REPETIR_ROTULOS = a.repetir_rotulos
    nomes = [n.strip() for n in a.abas.split(",") if n.strip()]
    desconhecidas = [n for n in nomes if n not in ABAS]
    if desconhecidas:
        sys.exit(f"abas desconhecidas: {desconhecidas}. Válidas: {list(ABAS)}")

    pasta = Path(a.brutos)
    if a.diagnostico:
        diagnostico(nomes, pasta)
        return
    pasta.mkdir(parents=True, exist_ok=True)

    resultado = {}
    for n in nomes:
        resultado[n] = extrair_aba(n, ABAS[n], pasta)
    alerta = gerar_municipios_alerta(resultado)
    if alerta is not None:
        resultado["municipiosAlerta"] = alerta

    validar(resultado)
    if a.csv_dir:
        gravar_csv(resultado, a.csv_dir)

    if a.modelo:
        modelo = Path(a.modelo)
        print("\nComparação com o modelo:")
        ref = pd.read_excel(modelo, sheet_name=None)
        for n, df in resultado.items():
            if n in ref:
                comparar_com_modelo(n, df, ref[n])
        if a.atualizar_modelo:
            print(f"\nArquivo atualizado: {atualizar_modelo(resultado, modelo)}")
            return

    if not a.sem_xlsx:
        gravar_xlsx(resultado, a.saida)
        print(f"\nArquivo gerado: {Path(a.saida).resolve()}")


if __name__ == "__main__":
    main()
